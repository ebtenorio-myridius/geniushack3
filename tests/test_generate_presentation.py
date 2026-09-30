import json

from pptx import Presentation

from tools import generate_presentation


def test_build_deck_when_latest_evaluation_has_no_successes(monkeypatch, tmp_path):
    results_path = tmp_path / "latest.json"
    results_path.write_text(
        json.dumps(
            {
                "case_count": 2,
                "results": [
                    {"case_id": "SYN-001", "success": False, "error_type": "APIConnectionError"},
                    {"case_id": "SYN-002", "success": False, "error_type": "APIConnectionError"},
                ],
            }
        ),
        encoding="utf-8",
    )
    output_path = tmp_path / "deck.pptx"
    monkeypatch.setattr(generate_presentation, "EVAL_RESULTS", results_path)
    monkeypatch.setattr(generate_presentation, "OUTPUT", output_path)
    monkeypatch.setattr(generate_presentation, "pytest_result", lambda: "60 passed")

    generate_presentation.build_deck()

    presentation = Presentation(output_path)
    slide_text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "0/2" in slide_text
    assert "No successful extraction cases" in slide_text
    assert "N/A" in slide_text