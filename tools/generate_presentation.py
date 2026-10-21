from pathlib import Path
import json
import re
import subprocess
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "presentation" / "Risk_Assessment_Workbench_Genius_Hacks_2026.pptx"
EVAL_RESULTS = ROOT / "evals" / "results" / "latest.json"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


# Palette: deep navy and royal blue with a restrained gold accent
NAVY = rgb("0A1B34")
NAVY_2 = rgb("132E55")
ROYAL = rgb("1F4E9C")
DECISION = rgb("163A7A")
MID = rgb("4A86D8")
SKY = rgb("9CC3F2")
ICE = rgb("EEF3FA")
LINE = rgb("D5DDEA")
INK = rgb("17253D")
MUTED = rgb("55657D")
GOLD = rgb("C8A24A")
GOLD_D = rgb("8A6A1F")
WHITE = rgb("FFFFFF")
SOFT = rgb("D6E4F7")  # secondary text on dark fills
SOFT_2 = rgb("E6EEF9")

HEAD_FONT = "Arial"
BODY_FONT = "Calibri"

# Canvas is 13.333 x 7.5 in (16:9); all coordinates below are in inches.
L = 0.7           # left margin
W = 11.93         # content width
FOOTER = "Risk Assessment Workbench  |  Genius Hacks 2026"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def evaluation_summary():
    data = json.loads(EVAL_RESULTS.read_text(encoding="utf-8"))
    successful = [item for item in data["results"] if item.get("success")]
    total_fields = sum(item.get("total_fields", 0) for item in successful)
    correct_fields = sum(item.get("correct_fields", 0) for item in successful)
    agreement_count = sum(bool(item.get("risk_agreement")) for item in successful)
    latency_values = [item["latency_ms"] for item in successful if "latency_ms" in item]
    return {
        "case_count": data.get("case_count", len(data["results"])),
        "successful_count": len(successful),
        "total_fields": total_fields,
        "correct_fields": correct_fields,
        "field_accuracy": correct_fields / total_fields if total_fields else None,
        "risk_agreement": agreement_count / len(successful) if successful else None,
        "mean_latency_ms": round(sum(latency_values) / len(latency_values)) if latency_values else None,
    }


def pytest_result():
    try:
        result = subprocess.run(
            [sys.executable, "-m", "pytest", "-q"],
            cwd=ROOT,
            capture_output=True,
            text=True,
            timeout=180,
            check=False,
        )
    except (OSError, subprocess.TimeoutExpired):
        return "not verified"

    output = f"{result.stdout}\n{result.stderr}"
    match = re.search(r"(\d+) passed(?:,|\s|$)", output)
    if result.returncode == 0 and match:
        return f"{match.group(1)} passed"
    if match:
        return f"{match.group(1)} passed; failures"
    return "not verified"


def per_case_accuracy():
    """Field accuracy (%) for each successful case, read from the latest live evaluation."""
    data = json.loads(EVAL_RESULTS.read_text(encoding="utf-8"))
    rows = []
    for item in data["results"]:
        if item.get("success") and item.get("total_fields"):
            rows.append((item["case_id"].split("-")[-1], round(100 * item["correct_fields"] / item["total_fields"])))
    return rows


# ---------------------------------------------------------------- helpers
def text_box(slide, text, x, y, w, h, size=16, color=INK, bold=False, font=BODY_FONT,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, spacing=None, tracking=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        if spacing:
            paragraph.space_after = Pt(spacing)
        run = paragraph.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        if tracking:
            run._r.get_or_add_rPr().set("spc", str(int(tracking * 100)))
    return box


def bullet_box(slide, items, x, y, w, h, size=16, color=INK, gap=10):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(gap)
        run = paragraph.add_run()
        run.text = item
        run.font.name = BODY_FONT
        run.font.size = Pt(size)
        run.font.color.rgb = color
        # real bullet with a hanging indent (no literal bullet characters in the text)
        pPr = paragraph._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(0.22))))
        pPr.set("indent", str(-int(Inches(0.22))))
        pPr.append(pPr.makeelement(f"{{{A_NS}}}buChar", {"char": "•"}))
    return box


def _add_shadow(item):
    """Soft navy drop shadow, equivalent to the cards in the designed deck."""
    sp_pr = item._element.spPr
    effect = sp_pr.makeelement(qn("a:effectLst"), {})
    shdw = effect.makeelement(qn("a:outerShdw"), {"blurRad": "177800", "dist": "38100", "dir": "5400000",
                                                  "algn": "t", "rotWithShape": "0"})
    clr = shdw.makeelement(qn("a:srgbClr"), {"val": "0A1B34"})
    clr.append(clr.makeelement(qn("a:alpha"), {"val": "12000"}))
    shdw.append(clr)
    effect.append(shdw)
    sp_pr.append(effect)


def _line_alpha(item, percent_opaque):
    clr = item._element.spPr.find(qn("a:ln")).find(qn("a:solidFill")).find(qn("a:srgbClr"))
    clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(percent_opaque * 1000))}))


def shape(slide, kind, x, y, w, h, fill, line=None, radius=0.08, line_w=0.75, shadow=False, name=None):
    item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        item.fill.background()
    else:
        item.fill.solid()
        item.fill.fore_color.rgb = fill
    if line is None:
        item.line.fill.background()
    else:
        item.line.color.rgb = line
        item.line.width = Pt(line_w)
    item.shadow.inherit = False
    for ref in item._element.iter(f"{{{A_NS}}}effectRef"):
        ref.set("idx", "0")  # no theme shadow in any renderer
    if shadow:
        _add_shadow(item)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        item.adjustments[0] = min(0.5, radius / min(w, h))
    if name:
        item.name = name
    return item


def rrect(slide, x, y, w, h, fill, line=None, shadow=False, radius=0.08, name=None):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, radius=radius, shadow=shadow, name=name)


def card(slide, x, y, w, h, fill=WHITE, line=LINE, shadow=True, name=None):
    return rrect(slide, x, y, w, h, fill, line, shadow=shadow, name=name)


def rule(slide, x, y, w, color=GOLD, width=1.5):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    connector.line.color.rgb = color
    connector.line.width = Pt(width)
    return connector


def badge(slide, number, x, y, d, fill):
    shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill)
    text_box(slide, str(number), x, y, d, d, round(d * 30), WHITE, True, HEAD_FONT, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def eyebrow(slide, text, dark=False):
    text_box(slide, text.upper(), L, 0.45, W, 0.3, 12, GOLD if dark else GOLD_D, True,
             anchor=MSO_ANCHOR.MIDDLE, tracking=3)


def new_slide(prs, state, kicker, title_text, dark=False, notes=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    state["n"] += 1
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY if dark else WHITE
    eyebrow(slide, kicker, dark)
    text_box(slide, title_text, L, 0.85, W, 0.8, 30, WHITE if dark else NAVY, True, HEAD_FONT)
    footer_color = SKY if dark else MUTED
    text_box(slide, FOOTER, L, 6.95, 8, 0.3, 10, footer_color, anchor=MSO_ANCHOR.MIDDLE)
    text_box(slide, str(state["n"]), 11.9, 6.95, 0.73, 0.3, 10, footer_color, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def chip(slide, x, y, w, h, label, fill, color, size=17, line=None):
    rrect(slide, x, y, w, h, fill, line, shadow=True)
    text_box(slide, label, x, y, w, h, size, color, True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def flow(slide, labels, styles, y, bw, h, size=17):
    gap = (W - len(labels) * bw) / (len(labels) - 1)
    for index, (label, (fill, color, line)) in enumerate(zip(labels, styles)):
        x = L + index * (bw + gap)
        chip(slide, x, y, bw, h, label, fill, color, size, line)
        if index < len(labels) - 1:
            text_box(slide, "→", x + bw, y, gap, h, 22, GOLD_D, False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def style_chart(chart, title_text, max_value):
    from pptx.enum.chart import XL_LABEL_POSITION

    chart.has_legend = False
    chart.has_title = True
    chart.chart_title.text_frame.text = title_text
    run = chart.chart_title.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(16)
    run.font.bold = False
    run.font.name = BODY_FONT
    run.font.color.rgb = NAVY
    chart.value_axis.visible = False
    chart.value_axis.has_major_gridlines = False
    chart.value_axis.maximum_scale = max_value
    chart.value_axis.minimum_scale = 0
    chart.category_axis.has_major_gridlines = False
    chart.category_axis.format.line.color.rgb = LINE
    chart.category_axis.tick_labels.font.size = Pt(14)
    chart.category_axis.tick_labels.font.name = BODY_FONT
    chart.category_axis.tick_labels.font.color.rgb = INK
    plot = chart.plots[0]
    plot.gap_width = 55
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.font.size = Pt(14)
    labels.font.name = BODY_FONT
    labels.font.color.rgb = INK
    labels.show_value = True
    labels.number_format = "0"
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END


# ---------------------------------------------------------------- deck
def build_deck():
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    evaluation = evaluation_summary()
    test_result = pytest_result()
    state = {"n": 0}
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1. Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    state["n"] += 1
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    for index, r in enumerate((2.0, 3.1, 4.2)):  # decorative rings, bottom right
        ring = shape(slide, MSO_SHAPE.OVAL, 12.2 - r, 7.9 - r, 2 * r, 2 * r, None, SKY, line_w=1)
        _line_alpha(ring, 20 + index * 4)
    text_box(slide, "GENIUS HACKS 2026", 0.9, 1.35, 6.5, 0.4, 16, GOLD, True, tracking=5)
    text_box(slide, "Risk Assessment Workbench", 0.9, 2.0, 6.6, 2.2, 48, WHITE, True, HEAD_FONT)
    text_box(slide, "From an unstructured change request to a governed, reviewable risk decision.", 0.9, 4.45, 6.2, 1.0, 22, SOFT)
    text_box(slide, "A synthetic-data vertical slice for financial crimes risk management", 0.9, 6.5, 7, 0.35, 14, SKY)
    panel = rrect(slide, 8.0, 1.2, 4.5, 5.0, NAVY_2, GOLD, radius=0.1)
    _line_alpha(panel, 45)
    for index, line_text in enumerate(["Prepare evidence.", "Preserve disagreement.", "Keep humans accountable."]):
        y = 1.45 + index * 1.5
        text_box(slide, line_text, 8.4, y, 3.7, 1.3, 26, WHITE, True, HEAD_FONT, anchor=MSO_ANCHOR.MIDDLE)
        if index < 2:
            rule(slide, 8.4, y + 1.4, 1.2)
    slide.notes_slide.notes_text_frame.text = "0:00-0:20. One-sentence pitch: the system prepares the evidence, humans make the decision. Introduce the team and roles."

    # 2. Problem
    slide = new_slide(prs, state, "The problem", "Slow, variable, and hard to reconstruct",
                      notes="0:20-0:50. Attribute the 15-20 business-day current-state figure to the brief and the two-day figure to the brief's target; neither is a measured app result. State the design stance: nothing is automatically approved or rejected.")
    y, h = 1.95, 4.65
    rrect(slide, L, y, 3.7, h, NAVY, shadow=True)
    text_box(slide, "15-20", L + 0.35, y + 0.4, 3.0, 1.1, 66, WHITE, True, HEAD_FONT, anchor=MSO_ANCHOR.MIDDLE)
    text_box(slide, "business days: brief estimate", L + 0.35, y + 1.55, 3.0, 0.4, 15, SOFT)
    rule(slide, L + 0.35, y + 2.4, 1.2)
    text_box(slide, "~2", L + 0.35, y + 2.65, 3.0, 1.1, 66, GOLD, True, HEAD_FONT, anchor=MSO_ANCHOR.MIDDLE)
    text_box(slide, "days: brief target, not measured", L + 0.35, y + 3.8, 3.0, 0.4, 15, SOFT)
    cw, gx = 3.8, 0.315
    x1 = L + 3.7 + 0.4
    x2 = x1 + cw + gx
    card(slide, x1, y, cw, h, ICE, LINE, shadow=False)
    text_box(slide, "Today", x1 + 0.35, y + 0.3, cw - 0.7, 0.5, 20, NAVY, True, HEAD_FONT)
    bullet_box(slide, ["Email, Word, Excel, SharePoint", "Different analysts reach different conclusions",
                       "Examiner asks why it was rated this way: evidence is scattered and slow to rebuild"],
               x1 + 0.35, y + 0.95, cw - 0.7, h - 1.25, 16, gap=14)
    card(slide, x2, y, cw, h, WHITE, ROYAL)
    text_box(slide, "Our response", x2 + 0.35, y + 0.3, cw - 0.7, 0.5, 20, ROYAL, True, HEAD_FONT)
    bullet_box(slide, ["One governed case record", "AI structures the evidence", "Humans own the decision",
                       "Every rating traceable to inputs and reasoning"],
               x2 + 0.35, y + 0.95, cw - 0.7, h - 1.25, 16, gap=14)

    # 3. Architecture
    slide = new_slide(prs, state, "Architecture", "A deliberately narrow AI boundary",
                      notes="0:50-1:50. Walk the pipeline left to right. Extraction is a probabilistic problem, scoring is not, so scoring is plain Python a committee can reproduce. This is the engineering-judgement answer.")
    plain = (ICE, INK, LINE)
    flow(slide, ["PDF", "pypdf", "LLM JSON", "Rules", "SQLite", "Review"],
         [plain, plain, (SKY, NAVY, None), (NAVY, WHITE, None), plain, (ROYAL, WHITE, None)], 2.0, 1.6, 0.85)
    for label, color, dx in [("Probabilistic (AI)", SKY, 0.0), ("Deterministic (code)", NAVY, 2.6), ("Human", ROYAL, 5.5)]:
        shape(slide, MSO_SHAPE.OVAL, L + dx, 3.16, 0.2, 0.2, color)
        text_box(slide, label, L + dx + 0.32, 3.1, 2.2, 0.32, 14, MUTED, anchor=MSO_ANCHOR.MIDDLE)
    cy, ch, cw2 = 3.75, 2.95, (W - 0.4) / 2
    for index, (name, fill, color, items) in enumerate([
        ("AI does", SKY, NAVY, ["Extract structure from unstructured text", "Identify stated risk factors", "Attempt semantic policy retrieval when configured"]),
        ("Code does", NAVY, WHITE, ["Validate schemas", "Calculate scores and thresholds", "Persist workflow events; policy lookup falls back to rules"]),
    ]):
        x = L + index * (cw2 + 0.4)
        card(slide, x, cy, cw2, ch)
        rrect(slide, x + 0.35, cy + 0.3, 1.9, 0.5, fill, radius=0.25)
        text_box(slide, name, x + 0.35, cy + 0.3, 1.9, 0.5, 16, color, True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        bullet_box(slide, items, x + 0.35, cy + 1.05, cw2 - 0.7, ch - 1.3, 16, gap=12)

    # 4. Harness
    slide = new_slide(prs, state, "AI harness and context", "The document is data, not authority",
                      notes="1:50-2:50. Context management: untrusted-document boundary, schema-constrained output, versioned prompts. If asked why no multi-agent orchestration: one narrow extraction step does not need it, and every extra agent is another probabilistic step to govern and evaluate.")
    columns = [
        ("Context boundary", ["Source text is explicitly untrusted", "Embedded instructions cannot approve a case", "Prompt requires faithful extraction"]),
        ("Structured contract", ["Pydantic schema at the model boundary", "Malformed output fails loudly", "Extraction versions stay reconstructable"]),
        ("Traceable context", ["Synthetic policy IDs and source paths", "Prompt version telemetry", "Analyst can disagree with the rationale"]),
    ]
    cw, y, h = (W - 0.8) / 3, 1.95, 4.75
    for index, (name, items) in enumerate(columns):
        x = L + index * (cw + 0.4)
        card(slide, x, y, cw, h)
        badge(slide, index + 1, x + 0.35, y + 0.35, 0.7, NAVY)
        text_box(slide, name, x + 0.35, y + 1.3, cw - 0.7, 0.5, 20, NAVY, True, HEAD_FONT)
        bullet_box(slide, items, x + 0.35, y + 2.0, cw - 0.7, h - 2.3, 16, gap=14)

    # 5. Scoring (native chart)
    slide = new_slide(prs, state, "Engineering judgement", "Scoring is explainable by construction",
                      notes="2:50-3:30. Deterministic versus probabilistic, justified. Be upfront that the weights are a prototype choice awaiting FCRM approval.")
    y, h, cwid = 1.95, 4.75, 6.6
    card(slide, L, y, cwid, h)
    data = CategoryChartData()
    data.categories = ["Customer + geography", "Product + channel", "Third party", "Complexity"]
    data.add_series("Weight (%)", (35, 30, 20, 15))
    frame = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(L + 0.2), Inches(y + 0.15),
                                   Inches(cwid - 0.4), Inches(h - 0.3), data)
    chart = frame.chart
    style_chart(chart, "Category weight in the risk score (%)", 40)
    chart.category_axis.reverse_order = True
    for index, color in enumerate([NAVY, ROYAL, MID, SKY]):
        point = chart.plots[0].series[0].points[index]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = color
    rx, rw, hh = L + cwid + 0.4, W - cwid - 0.4, (h - 0.4) / 2
    for index, (name, body, fill, color, bordered) in enumerate([
        ("Why not an LLM score?", "A committee must reproduce, challenge, and test the number. The model structures evidence; deterministic Python combines category scores.", ICE, NAVY, False),
        ("Honest limits", "Weights and policy material are synthetic, not approved supervisory mappings. Control effectiveness and residual risk are not yet modeled.", WHITE, GOLD_D, True),
    ]):
        cy = y + index * (hh + 0.4)
        card(slide, rx, cy, rw, hh, fill, GOLD if bordered else LINE, shadow=bordered)
        text_box(slide, name, rx + 0.35, cy + 0.25, rw - 0.7, 0.45, 20, color, True, HEAD_FONT)
        text_box(slide, body, rx + 0.35, cy + 0.8, rw - 0.7, hh - 1.0, 15, INK)

    # 6. Six stages
    slide = new_slide(prs, state, "Six stages", "AI is part of delivery, not only runtime code",
                      notes="3:30-4:15. One line per stage, naming the artifact: docs/requirements, docs/architecture, src and ai, tests and evals, ops. Point to docs/sdlc/ai-delivery-evidence.md. Move fast; the repo carries the detail.")
    stages = [("Requirements", "Brief → spec"), ("Design", "Options → decision"), ("Development", "Contract → code"),
              ("Testing", "Cases → metrics"), ("Deployment", "Image → runbook"), ("Operations", "Telemetry → feedback")]
    cw, ch, y0 = (W - 0.8) / 3, 1.85, 1.95
    for index, (name, detail) in enumerate(stages):
        x = L + (index % 3) * (cw + 0.4)
        y = y0 + (index // 3) * (ch + 0.3)
        card(slide, x, y, cw, ch)
        badge(slide, index + 1, x + 0.35, y + 0.3, 0.65, ROYAL if index % 2 else NAVY)
        text_box(slide, name, x + 1.2, y + 0.3, cw - 1.5, 0.65, 20, NAVY, True, HEAD_FONT, anchor=MSO_ANCHOR.MIDDLE)
        text_box(slide, detail, x + 0.35, y + 1.2, cw - 0.7, 0.4, 17, MUTED, anchor=MSO_ANCHOR.MIDDLE)
    by = y0 + 2 * ch + 0.6
    rrect(slide, L, by, W, 0.65, NAVY)
    text_box(slide, "Each stage has an AI instruction, a human gate, and a repository artifact.", L + 0.35, by, W - 0.7, 0.65,
             18, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)

    # 7. Governance
    slide = new_slide(prs, state, "Human-in-the-loop", "The system prepares; people decide",
                      notes="4:15-5:15. Every risk level goes through analyst review before committee decision. Explain the resume-after-information-request path and that conditions must be evidenced before analyst closure. The event hash chain and examiner export help reconstruct the record but are not immutable against a database owner.")
    flow(slide, ["Owner draft", "Analyst review", "Committee", "Decision", "Close"],
         [(SKY, NAVY, None), (MID, WHITE, None), (ROYAL, WHITE, None), (DECISION, WHITE, None), (NAVY, WHITE, None)],
         1.95, 2.0, 0.75)
    gates = [
        ("Analyst gate", ["Every risk level reviewed", "Edit/rescore with source citations", "Finalize with rationale"]),
        ("Committee + lifecycle", ["Three distinct votes; 2-vote majority", "Information requests resume prior stage", "Conditions block closure; reassessments link"]),
        ("Audit evidence", ["Versioned extractions and UTC events", "Verifiable event hash chain + source hashes", "Examiner ZIP with verification manifest"]),
    ]
    cw3, cy, ch = (W - 0.8) / 3, 3.05, 3.2
    for index, (name, items) in enumerate(gates):
        x = L + index * (cw3 + 0.4)
        card(slide, x, cy, cw3, ch)
        badge(slide, index + 1, x + 0.3, cy + 0.28, 0.55, NAVY)
        text_box(slide, name, x + 1.05, cy + 0.28, cw3 - 1.3, 0.55, 18, NAVY, True, HEAD_FONT, anchor=MSO_ANCHOR.MIDDLE)
        bullet_box(slide, items, x + 0.3, cy + 1.1, cw3 - 0.6, ch - 1.3, 15, gap=10)
    text_box(slide, "Demo identities only; vote/scoring rules are unapproved. Hash-chain verification is not immutable storage.",
             L, 6.4, W, 0.35, 13, MUTED, italic=True, anchor=MSO_ANCHOR.MIDDLE)

    # 8. Demo
    slide = new_slide(prs, state, "Working demonstration", "One synthetic case, end to end",
                      notes="5:15-7:45 (about 2.5 minutes). Owner uploads, checks required draft fields, and submits. The analyst verifies/corrects with source citations, finalizes, and submits to committee. Use three separate committee identities. For a conditional outcome, show owner evidence followed by analyst closure; end with audit verification/export. If the provider is unavailable, use the recorded evaluation and local tests without implying offline extraction.")
    steps = ["Owner uploads PDF", "Inspect draft + score", "Owner submits case", "Analyst finalizes", "Three committee votes", "Owner evidence; analyst closes"]
    colw, d, cy = W / 6, 0.8, 2.1
    rule(slide, L + colw / 2, cy + d / 2, W - colw, SKY, 3)
    for index, step in enumerate(steps):
        x = L + index * colw
        badge(slide, index + 1, x + colw / 2 - d / 2, cy, d, NAVY if index < 3 else ROYAL)
        text_box(slide, step, x + 0.1, cy + d + 0.2, colw - 0.2, 0.9, 15, INK, True, align=PP_ALIGN.CENTER)
    by, bh = 4.55, 2.1
    card(slide, L, by, W, bh, NAVY, None)
    text_box(slide, "Security boundary", L + 0.5, by + 0.3, W - 1, 0.5, 20, GOLD, True, HEAD_FONT)
    text_box(slide, "Workflow routes check fixed demo roles, not production identities. Extraction needs the configured model provider; no offline extraction fallback. Source and audit hashes are verifiable, not immutable.",
             L + 0.5, by + 0.9, W - 1, 1.0, 17, WHITE)

    # 9. Evaluation (numbers come from evals/results/latest.json)
    slide = new_slide(prs, state, "Testing and evidence", "We measure the model, not just the demo",
                      notes="7:45-9:45. Quote the values generated from the latest recorded PDF evaluation and state that this is only eight synthetic cases. Risk agreement comes from deterministic scoring and does not establish field accuracy. Mention controlled errors for invalid PDFs and provider failures; do not suggest extraction has an offline fallback.")
    field_accuracy = f"{evaluation['field_accuracy']:.1%}" if evaluation["field_accuracy"] is not None else "N/A"
    risk_agreement = f"{evaluation['risk_agreement']:.0%}" if evaluation["risk_agreement"] is not None else "N/A"
    latency = f"{evaluation['mean_latency_ms'] / 1000:.1f} s" if evaluation["mean_latency_ms"] is not None else "N/A"
    tiles = [(f"{evaluation['successful_count']}/{evaluation['case_count']}", "live calls succeeded"),
             (field_accuracy, f"field accuracy ({evaluation['correct_fields']}/{evaluation['total_fields']})"),
             (risk_agreement, "risk-level agreement"), (latency, "mean latency per case")]
    tw, ty, th = (W - 0.9) / 4, 1.85, 1.3
    for index, (value, label) in enumerate(tiles):
        x = L + index * (tw + 0.3)
        rrect(slide, x, ty, tw, th, ROYAL if index % 2 else NAVY, shadow=True)
        text_box(slide, value, x, ty + 0.1, tw, 0.75, 38, WHITE, True, HEAD_FONT, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        text_box(slide, label, x, ty + 0.85, tw, 0.35, 14, SOFT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    y, h, cwid = 3.4, 3.35, 6.0
    rows = per_case_accuracy()
    card(slide, L, y, cwid, h)
    if rows:
        chart_data = CategoryChartData()
        chart_data.categories = [case for case, _ in rows]
        chart_data.add_series("Field accuracy (%)", [value for _, value in rows])
        frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(L + 0.15), Inches(y + 0.1),
                                       Inches(cwid - 0.3), Inches(h - 0.2), chart_data)
        style_chart(frame.chart, "Field accuracy by synthetic case (%)", 125)
        series = frame.chart.plots[0].series[0]
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = ROYAL
    else:
        text_box(slide, "No successful extraction cases", L + 0.35, y + 0.3, cwid - 0.7, 0.5, 20, NAVY, True, HEAD_FONT)
        text_box(slide, "The latest evaluation could not measure field accuracy. Show the recorded provider errors and rerun when the model endpoint is reachable.",
                 L + 0.35, y + 0.95, cwid - 0.7, 1.6, 16, INK)
    rx, rw = L + cwid + 0.4, W - cwid - 0.4
    card(slide, rx, y, rw, h, ICE, LINE, shadow=False)
    text_box(slide, "Failure handled, and how to read this", rx + 0.35, y + 0.22, rw - 0.7, 0.45, 17, NAVY, True, HEAD_FONT)
    bullet_box(slide, [
        "Invalid PDFs and provider failures return explicit errors; transient provider failures are retried up to three times",
        "Risk agreement comes from the deterministic scorer; it does not prove every field is right",
        f"Eight synthetic cases, not a production accuracy estimate. pytest: {test_result}",
    ], rx + 0.35, y + 0.8, rw - 0.7, h - 0.95, 14, gap=8)

    # 10. Close
    slide = new_slide(prs, state, "Deployment and next steps", "Prototype-ready, with an honest path to production", dark=True,
                      notes="9:45-10:00. Close on the tagline, then stop on time. Leave the five minutes of Q&A for the decisions: why deterministic scoring, why no multi-agent orchestration, why field accuracy and risk agreement differ.")
    columns = [
        ("Current demo", ["Docker Compose", "SQLite named volume", "Synthetic data only", "OpenAI key for extraction"]),
        ("Free-first target", ["Ollama adapter", "PostgreSQL and MinIO", "Redis worker", "Prometheus and Grafana"]),
        ("Production gate", ["Real identity + hardened access", "Approved rules + residual-risk controls", "Protected storage + external audit anchor", "Migrations, backups, alerts, rollback"]),
    ]
    cw, y, h = (W - 0.8) / 3, 1.95, 3.95
    for index, (name, items) in enumerate(columns):
        x = L + index * (cw + 0.4)
        box = rrect(slide, x, y, cw, h, NAVY_2, GOLD if index == 2 else SKY)
        _line_alpha(box, 70 if index == 2 else 30)
        text_box(slide, name, x + 0.35, y + 0.28, cw - 0.7, 0.5, 20, GOLD if index == 2 else WHITE, True, HEAD_FONT)
        bullet_box(slide, items, x + 0.35, y + 0.95, cw - 0.7, h - 1.15, 16, SOFT_2, gap=11)
    text_box(slide, "Prepare evidence. Preserve disagreement. Keep humans accountable.", L, 6.05, W, 0.7, 22, GOLD, True,
             HEAD_FONT, anchor=MSO_ANCHOR.MIDDLE)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_deck()
